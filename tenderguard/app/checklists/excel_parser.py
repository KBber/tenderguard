"""Excel -> normalized TenderGuard v0.3 checklist (CK with atomic requirements)."""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

try:
    import openpyxl
except ImportError as exc:
    raise SystemExit("openpyxl is required: pip install openpyxl") from exc

from tenderguard.app.checklists.review_qa import derive_review_qa
from tenderguard.app.schemas import (
    AtomicRequirement,
    AtomicVerification,
    ChecklistItem,
    EvidenceContract,
    EvidenceQuality,
    Precondition,
    ReviewReason,
    ReviewType,
    Severity,
    VerificationMethod,
)


SHEETS_TO_PROCESS = ("直投CK", "非直投CK", "列投标总体CK")


# Title rules (spec §十二)
_TITLES: dict[str, str] = {
    "招标文件获取": "招标文件获取",
    "招标信息检查": "采购人主体信息核查",
    "招标文件内容分析": "招标文件内容核对",
    "补充材料": "招标补充材料跟踪",
    "标前任务": "标前任务规划",
    "分析招标文件": "招标文件风险标注",
    "投标策略确认": "投标策略确认",
    "任务分工": "投标任务分工",
    "演示设备借用": "演示设备借用",
    "项目启动会": "项目启动会",
    "项目夕会": "项目夕会",
    "风险点评审": "投标风险点评审",
    "模拟打分": "投标模拟打分",
    "项目汇报": "投标项目日报",
    "项目信息闭环": "项目信息闭环",
    "投标报名": "投标报名",
    "投标人": "投标人直投主体",
    "模板与格式": "标书模板与格式",
    "是否标书结构中囊括了所有招标要求": "标书结构完整性",
    "保证金": "投标保证金申请",
    "授权": "授权文件准备",
    "投标RP制作": "投标RP制作",
    "商务偏离表": "商务偏离表核心响应",
    "分项报价表": "分项报价表响应",
    "本国产品标准证明文件": "本国产品标准证明",
    "投标方案": "投标方案引用合规",
    "投标清单": "投标清单设备选型",
    "失信名单排查": "军队采购失信名单排查",
    "节能环保": "节能环保产品合规",
    "投标报价": "投标报价决策",
    "资质材料": "资质材料完整性",
    "产品资质": "产品资质完整性",
    "产品参数": "投标产品参数偏高核查",
    "公司更名证明": "公司更名证明",
    "业绩": "业绩合规性",
    "公章与法人章": "公章与法人章",
    "整体检查": "整体合规性检查",
    "一致性检查": "跨文档一致性检查",
    "投标保证金内容检查": "投标保证金内容结构化校验",
    "三级审查机制": "三级审查机制",
    "打印、封标": "打印与封装",
    "应答、签到、解密": "电子应答与解密",
    "述标": "讲标准备",
    "现场投标，签字确认等": "现场投标与签字",
    "样品与测试环境": "样品与测试环境",
    "复盘": "投标项目复盘",
}


def _clean_title(name: str) -> str:
    if not name:
        return "(未命名)"
    m = re.match(r"^(\d+)[、.]", name)
    if m:
        name = name.split("、", 1)[-1].split(".", 1)[-1].strip()
    for key, title in _TITLES.items():
        if key in name:
            return title
    return name[:24]


# Classifier keywords
_PROCESS_KEYWORDS = (
    "召开", "夕会", "启动会", "日报", "周报", "邮件", "汇报", "互审", "复盘",
    "现场投标", "讲标", "述标", "签到", "解密", "封装",
    "运输", "搬运", "演示借用", "三级审查",
)
_STRATEGY_KEYWORDS = (
    "投标策略", "商业策略", "竞争对手报价", "商务处理", "策略确认",
    "是否为我方", "是否支持", "模拟打分", "评分表", "商务策略",
    "竞争对手的报价", "销售经理",
)
_EXTERNAL_KEYWORDS = (
    "军队采购失信名单", "失信名单", "查询", "官网", "政府采购",
    "节能产品政府采购品目", "环境标志产品政府采购品目",
)
_PROJECT_META = ("项目名称", "项目编号", "包号", "投标人", "法人代表", "被授权人", "营业执照")
_BOND = ("保证金", "投标保证金")
_PRICING = ("分项报价表", "型号不同报价", "同型号", "限价", "拦标价", "单价", "小计", "合计", "报价", "出货清单", "财评清单", "设备清单")
_RESPONSE = ("商务偏离表", "商务响应", "技术响应", "完整响应", "废标项", "评分项", "逐条", "应答", "偏离", "商务条件", "商务条款", "商务要求")
_CERTIFICATE = ("资质材料", "产品资质", "检测报告", "节能证书", "3C", "产品彩页", "本国产品标准", "公章", "盖章", "法人章", "CCC", "彩页")
_PARAM = ("产品参数", "设备选型", "投标清单", "技术参数", "参数", "符合招标", "技术响应")
_EXPERIENCE = ("业绩", "中标通知书", "验收报告")
_CONSISTENCY = ("一致性", "与招标文件一致", "项目名称", "项目编号", "包号")
_COVERAGE_FULL_DOC = ("标书结构", "整体检查", "整体响应", "完全响应", "全部响应", "是否标书结构中")


def _classify_v3(category: str, name: str, description: str, row_idx: int) -> dict[str, Any]:
    text = f"\n{category}\n{name}\n{description}"

    if any(k in text for k in _PROCESS_KEYWORDS):
        return _process_human()
    if any(k in text for k in _STRATEGY_KEYWORDS):
        return _strategy_human()
    if any(k in text for k in _EXTERNAL_KEYWORDS):
        return _external_data()
    if any(k in text for k in _COVERAGE_FULL_DOC):
        return _coverage_semantic()

    if any(k in text for k in _CERTIFICATE):
        return _certificate_hybrid()
    if any(k in text for k in _RESPONSE):
        return _response_hybrid()
    if any(k in text for k in _PRICING):
        # Exclude only process-as-pricing contexts (engineering output / updates).
        if "工程部输出" in text or "更新过设备报价" in text or "询价时必须" in text:
            return _process_human()
        return _pricing_deterministic()
    if any(k in text for k in _PROJECT_META):
        return _project_meta_deterministic()
    if any(k in text for k in _CONSISTENCY):
        return _consistency_deterministic()
    if "投标报价" in text and ("授权" in text or "审核" in text):
        return _strategy_human()
    if any(k in text for k in _BOND):
        return _bond_hybrid()

    if any(k in text for k in _PARAM):
        return _param_hybrid()
    if any(k in text for k in _EXPERIENCE):
        return _experience_hybrid()

    return _process_human()


def _req(req_id, desc, source, verif, spec=None):
    from tenderguard.app.schemas import AtomicRequirement
    return AtomicRequirement(
        requirement_id=req_id, description=desc, source=source,
        verification=verif, spec=spec or {},
    )


def _ev_contract(min_q=EvidenceQuality.DIRECT, docs=None):
    from tenderguard.app.schemas import EvidenceContract
    return EvidenceContract(minimum_quality=min_q, required_documents=docs or ["tender", "bid"])


def _process_human():
    return {
        "review_type": ReviewType.PROCESS_HUMAN,
        "classification_reason": "CK 描述的是企业内部流程动作（启动会、夕会、复盘、封装、述标等），无法仅通过招标文件 + 投标文件判定。",
        "atomic_requirements": [], "required_evidence": [], "rule": None,
        "evidence_contract": _ev_contract(EvidenceQuality.MISSING, []),
        "preconditions": [],
        "failure_taxonomy": [ReviewReason.HUMAN_PROCESS_REQUIRED],
    }


def _strategy_human():
    return {
        "review_type": ReviewType.STRATEGY_HUMAN,
        "classification_reason": "CK 涉及商业策略判断（投标策略、竞争对手报价、模拟打分），属业务决策层。",
        "atomic_requirements": [], "required_evidence": [], "rule": None,
        "evidence_contract": _ev_contract(EvidenceQuality.MISSING, []),
        "preconditions": [],
        "failure_taxonomy": [ReviewReason.STRATEGY_REQUIRED],
    }


def _external_data():
    return {
        "review_type": ReviewType.EXTERNAL_DATA,
        "classification_reason": "CK 需查询外部权威系统（军队采购失信名单、政府采购目录、产品官网）。",
        "atomic_requirements": [], "required_evidence": [], "rule": None,
        "evidence_contract": _ev_contract(EvidenceQuality.MISSING, []),
        "preconditions": [],
        "failure_taxonomy": [ReviewReason.EXTERNAL_DATA_REQUIRED],
    }


def _pricing_deterministic():
    return {
        "review_type": ReviewType.DOCUMENT_DETERMINISTIC,
        "classification_reason": "价格 / 同型号 / 小计 / 限价 均为代码可确定，无需 LLM。",
        "atomic_requirements": [
            _req("PRICE-001", "同型号产品报价必须一致", "投标文件", AtomicVerification.SAME_MODEL_SAME_PRICE, {"operator": "same_model_same_price", "key": "model", "value": "unit_price"}),
            _req("PRICE-002", "每个报价行小计 = 数量 × 单价", "投标文件", AtomicVerification.TABLE_SUM, {"operator": "sum_equals", "rows": "bid.price_rows"}),
            _req("PRICE-003", "投标总价不得超过限价", "跨文档", AtomicVerification.NUMERIC_COMPARISON, {"operator": "numeric_lte", "actual_field": "bid.total_price", "required_field": "tender.price_ceiling"}),
        ],
        "required_evidence": ["bid.price_rows", "bid.total_price", "tender.price_ceiling"],
        "rule": {"operator": "same_model_same_price", "key": "model", "value": "unit_price"},
        "evidence_contract": _ev_contract(EvidenceQuality.DIRECT, ["bid", "tender"]),
        "preconditions": [
            Precondition(expression="bid.price_rows is not empty", description="报价表存在"),
            Precondition(expression="tender.price_ceiling is not None", description="限价已识别"),
        ],
        "failure_taxonomy": [ReviewReason.MISSING_EVIDENCE, ReviewReason.TABLE_EXTRACTION_FAILED, ReviewReason.EVIDENCE_CONFLICT],
    }


def _project_meta_deterministic():
    return {
        "review_type": ReviewType.DOCUMENT_DETERMINISTIC,
        "classification_reason": "项目名称/编号/包号/投标人/法人代表等元数据可逐字段等价比对，代码可判定。",
        "atomic_requirements": [
            _req("META-001", "项目名称跨文档一致", "跨文档", AtomicVerification.FIELD_CONSISTENCY, {"operator": "equals", "fields": ["tender.project_name", "bid.project_name"]}),
            _req("META-002", "项目编号跨文档一致", "跨文档", AtomicVerification.FIELD_CONSISTENCY, {"operator": "equals", "fields": ["tender.project_id", "bid.project_id"]}),
            _req("META-003", "包号跨文档一致", "跨文档", AtomicVerification.FIELD_CONSISTENCY, {"operator": "equals", "fields": ["tender.package_id", "bid.package_id"]}),
            _req("META-004", "投标人名称与营业执照一致", "跨文档", AtomicVerification.FIELD_CONSISTENCY, {"operator": "field_consistency", "pairs": [["bid.bidder_name", "license.company_name"]]}),
            _req("META-005", "法人代表 / 被授权人一致", "投标文件", AtomicVerification.FIELD_CONSISTENCY, {"operator": "field_consistency", "pairs": [["authorization.legal_representative", "bid.legal_representative"], ["authorization.authorized_person", "authorization.legal_representative"]]}),
        ],
        "required_evidence": ["tender.project_name", "bid.project_name", "tender.project_id", "bid.project_id", "tender.package_id", "bid.package_id", "bid.bidder_name", "license.company_name", "authorization.legal_representative", "bid.legal_representative", "authorization.authorized_person"],
        "rule": {"operator": "equals", "fields": ["tender.project_name", "bid.project_name"]},
        "evidence_contract": _ev_contract(EvidenceQuality.DIRECT, ["tender", "bid"]),
        "preconditions": [
            Precondition(expression="tender.project_name is not None", description="招标项目名称已抽取"),
            Precondition(expression="bid.project_name is not None", description="投标项目名称已抽取"),
        ],
        "failure_taxonomy": [ReviewReason.EXTRACTION_FAILED, ReviewReason.MISSING_EVIDENCE, ReviewReason.EVIDENCE_CONFLICT],
    }


def _consistency_deterministic():
    return {
        "review_type": ReviewType.DOCUMENT_DETERMINISTIC,
        "classification_reason": "跨文档一致性由 equals / field_consistency 算子直接判定。",
        "atomic_requirements": [
            _req("CONS-001", "证书/资质编号跨表格一致", "投标文件", AtomicVerification.FIELD_CONSISTENCY, {"operator": "field_consistency", "pairs": [["deviation.test_report_no", "report.number"], ["deviation.test_report_name", "report.name"]]}),
            _req("CONS-002", "合同金额与业绩汇总一致", "投标文件", AtomicVerification.FIELD_CONSISTENCY, {"operator": "equals", "fields": ["experience.contract_amount", "experience.summary_amount"]}),
            _req("CONS-003", "项目编号/名称/包号一致", "跨文档", AtomicVerification.FIELD_CONSISTENCY, {"operator": "equals", "fields": ["tender.project_id", "bid.project_id", "tender.package_id", "bid.package_id"]}),
            _req("CONS-004", "同型号同价", "投标文件", AtomicVerification.SAME_MODEL_SAME_PRICE, {"operator": "same_model_same_price"}),
        ],
        "required_evidence": ["tender.project_id", "bid.project_id", "tender.package_id", "bid.package_id", "deviation.test_report_no", "report.number", "bid.price_rows"],
        "rule": {"operator": "equals", "fields": ["tender.project_id", "bid.project_id"]},
        "evidence_contract": _ev_contract(EvidenceQuality.DIRECT, ["tender", "bid"]),
        "preconditions": [
            Precondition(expression="tender.project_id is not None", description="招标编号已抽取"),
            Precondition(expression="bid.project_id is not None", description="投标编号已抽取"),
        ],
        "failure_taxonomy": [ReviewReason.EXTRACTION_FAILED, ReviewReason.EVIDENCE_CONFLICT],
    }


def _bond_hybrid():
    return {
        "review_type": ReviewType.DOCUMENT_HYBRID,
        "classification_reason": "保证金金额 + 收款方 + 截止时间 需抽取结构化事实后由 Rule 判定。",
        "atomic_requirements": [
            _req("BOND-A", "保证金金额满足招标文件要求", "跨文档", AtomicVerification.NUMERIC_COMPARISON, {"operator": "numeric_gte", "actual_field": "bond.amount", "required_field": "requirement.bond_amount"}),
            _req("BOND-B", "收款方与招标一致", "跨文档", AtomicVerification.FIELD_CONSISTENCY, {"operator": "equals", "fields": ["bond.payee", "requirement.bond_payee"]}),
            _req("BOND-C", "截止时间满足招标要求", "跨文档", AtomicVerification.DATE_RANGE, {"operator": "date_before", "actual_field": "bond.deadline", "required_field": "requirement.bond_deadline"}),
        ],
        "required_evidence": ["bond.amount", "bond.payee", "bond.deadline", "requirement.bond_amount", "requirement.bond_payee", "requirement.bond_deadline"],
        "rule": {"operator": "field_consistency", "pairs": [["bond.payee", "requirement.bond_payee"]]},
        "evidence_contract": _ev_contract(EvidenceQuality.DIRECT, ["tender", "bid"]),
        "preconditions": [
            Precondition(expression="bond.amount is not None", description="投标保证金金额已抽取"),
            Precondition(expression="bond.payee is not None", description="收款方已抽取"),
        ],
        "failure_taxonomy": [ReviewReason.REQUIREMENT_UNRESOLVED, ReviewReason.MISSING_EVIDENCE, ReviewReason.EVIDENCE_CONFLICT],
    }


def _certificate_hybrid():
    return {
        "review_type": ReviewType.DOCUMENT_HYBRID,
        "classification_reason": "资质覆盖 = 招标要求列表 vs 投标提供列表（COVERAGE 算子 + 0.8 阈值）。",
        "atomic_requirements": [
            _req("CERT-001", "产品资质覆盖 ≥ 0.8", "跨文档", AtomicVerification.COVERAGE, {"operator": "coverage", "source": "tender.product_cert_requirements", "target": "bid.product_certificates", "min_coverage": 0.8}),
            _req("CERT-002", "资质材料在有效期内", "投标文件", AtomicVerification.FIELD_CONSISTENCY, {"operator": "required", "field": "bid.certificates_validity"}),
        ],
        "required_evidence": ["tender.product_cert_requirements", "bid.product_certificates"],
        "rule": {"operator": "coverage", "source": "tender.product_cert_requirements", "target": "bid.product_certificates", "min_coverage": 0.8},
        "evidence_contract": _ev_contract(EvidenceQuality.SUPPORTING, ["tender", "bid"]),
        "preconditions": [
            Precondition(expression="tender.product_cert_requirements is not empty", description="招标资质要求已抽取"),
        ],
        "failure_taxonomy": [ReviewReason.MISSING_EVIDENCE, ReviewReason.EXTRACTION_FAILED],
    }


def _param_hybrid():
    return {
        "review_type": ReviewType.DOCUMENT_HYBRID,
        "classification_reason": "产品参数响应 = 抽取参数列表 → 比对 bid_meets_or_exceeds。",
        "atomic_requirements": [
            _req("PARAM-001", "投标参数满足或超过招标参数", "跨文档", AtomicVerification.SEMANTIC_MATCH, {"operator": "parameter_compare", "direction": "bid_meets_or_exceeds"}),
            _req("PARAM-002", "投标型号在官网可追溯", "投标文件", AtomicVerification.FIELD_CONSISTENCY, {"operator": "required", "field": "bid.product_traces"}),
        ],
        "required_evidence": ["tender.tech_params", "bid.tech_params", "bid.product_traces"],
        "rule": {"operator": "parameter_compare", "direction": "bid_meets_or_exceeds"},
        "evidence_contract": _ev_contract(EvidenceQuality.SUPPORTING, ["tender", "bid"]),
        "preconditions": [Precondition(expression="tender.tech_params is not empty", description="招标参数已抽取")],
        "failure_taxonomy": [ReviewReason.MISSING_EVIDENCE, ReviewReason.EXTRACTION_FAILED],
    }


def _experience_hybrid():
    return {
        "review_type": ReviewType.DOCUMENT_HYBRID,
        "classification_reason": "业绩需 LLM 抽取合同金额/时间，再由 Rule 验证数量与时间窗。",
        "atomic_requirements": [
            _req("EXP-001", "业绩数量满足要求", "投标文件", AtomicVerification.NUMERIC_COMPARISON, {"operator": "count_gte", "rows": "bid.experiences", "count_min": 3}),
            _req("EXP-002", "业绩合同金额满足门槛", "跨文档", AtomicVerification.NUMERIC_COMPARISON, {"operator": "numeric_gte", "actual_field": "experience.contract_amount", "required_field": "requirement.experience_contract_amount"}),
            _req("EXP-003", "业绩在招标时间窗内", "跨文档", AtomicVerification.DATE_RANGE, {"operator": "date_range", "rows": "bid.experiences"}),
        ],
        "required_evidence": ["bid.experiences", "requirement.experience_contract_amount", "requirement.experience_period_years"],
        "rule": {"operator": "count_gte", "rows": "bid.experiences", "count_min": 3},
        "evidence_contract": _ev_contract(EvidenceQuality.SUPPORTING, ["bid"]),
        "preconditions": [Precondition(expression="bid.experiences is not empty", description="业绩列表已抽取")],
        "failure_taxonomy": [ReviewReason.REQUIREMENT_UNRESOLVED, ReviewReason.MISSING_EVIDENCE, ReviewReason.EXTRACTION_FAILED],
    }


def _response_hybrid():
    return {
        "review_type": ReviewType.DOCUMENT_HYBRID,
        "classification_reason": "商务/技术响应需 LLM 抽取商务条件条目，再做 coverage 比对。",
        "atomic_requirements": [
            _req("RESP-001", "商务要求覆盖 ≥ 1.0", "跨文档", AtomicVerification.COVERAGE, {"operator": "coverage", "source": "tender.commercial_requirements", "target": "bid.commercial_responses", "min_coverage": 1.0}),
            _req("RESP-002", "技术要求覆盖 ≥ 1.0", "跨文档", AtomicVerification.COVERAGE, {"operator": "coverage", "source": "tender.technical_requirements", "target": "bid.technical_responses", "min_coverage": 1.0}),
        ],
        "required_evidence": ["tender.commercial_requirements", "bid.commercial_responses", "tender.technical_requirements", "bid.technical_responses"],
        "rule": {"operator": "coverage", "source": "tender.commercial_requirements", "target": "bid.commercial_responses", "min_coverage": 1.0},
        "evidence_contract": _ev_contract(EvidenceQuality.SUPPORTING, ["tender", "bid"]),
        "preconditions": [Precondition(expression="tender.commercial_requirements is not empty", description="商务要求已抽取")],
        "failure_taxonomy": [ReviewReason.MISSING_EVIDENCE, ReviewReason.EXTRACTION_FAILED],
    }


def _coverage_semantic():
    return {
        "review_type": ReviewType.DOCUMENT_SEMANTIC,
        "classification_reason": "需 LLM 理解文档结构（章节 / 子章节）后做 section-level coverage。",
        "atomic_requirements": [
            _req("DOC-001", "标书结构覆盖全部招标要求章节", "跨文档", AtomicVerification.SEMANTIC_MATCH, {"operator": "coverage", "source": "tender.required_sections", "target": "bid.submitted_sections", "min_coverage": 1.0}),
            _req("DOC-002", "技术方案无负偏离", "投标文件", AtomicVerification.SEMANTIC_MATCH, {"operator": "no_negative_deviation"}),
        ],
        "required_evidence": ["tender.required_sections", "bid.submitted_sections"],
        "rule": {"operator": "coverage", "source": "tender.required_sections", "target": "bid.submitted_sections", "min_coverage": 1.0},
        "evidence_contract": _ev_contract(EvidenceQuality.INDIRECT, ["tender", "bid"]),
        "preconditions": [Precondition(expression="tender.required_sections is not empty", description="招标章节要求已抽取")],
        "failure_taxonomy": [ReviewReason.REQUIREMENT_UNRESOLVED, ReviewReason.MISSING_EVIDENCE],
    }


def _severity_for(category: str, name: str) -> Severity:
    text = category + name
    if any(k in text for k in ("废标", "保证金", "报价", "项目名称", "项目编号", "包号", "资质", "印章", "授权", "法人")):
        return Severity.CRITICAL
    if any(k in text for k in ("业绩", "偏离", "偏离表", "技术参数", "产品参数", "响应", "覆盖率", "检测报告", "节能", "3C", "彩页")):
        return Severity.HIGH
    if "复盘" in text or "总结" in text:
        return Severity.LOW
    return Severity.MEDIUM


def parse_excel(path: str | Path, sheets: tuple[str, ...] = SHEETS_TO_PROCESS) -> list[ChecklistItem]:
    wb = openpyxl.load_workbook(path, data_only=True)
    items: list[ChecklistItem] = []
    counter = 0
    for sheet_name in sheets:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        current_category: str | None = None
        for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            values = [str(c).strip() if c is not None else "" for c in row]
            if row_idx <= 2:
                continue
            if all(not v for v in values):
                continue
            if values[:5] == ["大类", "检查项", "检查内容", "检查明细", "执行结果"]:
                continue
            big_cat, item_name, content, detail, _result = (values + [""] * 5)[:5]
            if big_cat:
                current_category = big_cat
            category = current_category or "未分类"
            counter += 1
            check_id = f"DIRECT-{counter:03d}"
            description = "\n".join(p for p in (content, detail) if p and p.strip())
            cls = _classify_v3(category, item_name or "", description, row_idx)
            title = _clean_title(item_name or description[:30])
            severity = _severity_for(category, item_name or "")
            review_type = cls["review_type"]
            q, a = derive_review_qa(title, review_type)
            item = ChecklistItem(
                check_id=check_id,
                title=title,
                source_text=description or item_name or "",
                category=category,
                severity=severity,
                review_type=review_type,
                classification_reason=cls["classification_reason"],
                atomic_requirements=cls["atomic_requirements"],
                required_evidence=cls["required_evidence"],
                rule=cls["rule"],
                evidence_contract=cls["evidence_contract"],
                preconditions=cls["preconditions"],
                verification_method=VerificationMethod.DETERMINISTIC if review_type == ReviewType.DOCUMENT_DETERMINISTIC
                                   else (VerificationMethod.HYBRID if review_type == ReviewType.DOCUMENT_HYBRID else
                                         (VerificationMethod.SEMANTIC if review_type == ReviewType.DOCUMENT_SEMANTIC else
                                          VerificationMethod.HUMAN_REVIEW)),
                failure_taxonomy=cls["failure_taxonomy"],
                review_question=q,
                recommended_action=a,
                source_sheet=sheet_name,
                source_row=row_idx,
                original_requirement=item_name or "",
                notes="",
            )
            items.append(item)
    return items


def parse_to_file(excel_path, output_path, sheets: tuple[str, ...] = SHEETS_TO_PROCESS) -> int:
    items = parse_excel(excel_path, sheets)
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "source_excel": str(excel_path),
        "sheets": list(sheets),
        "count": len(items),
        "version": "v0.3",
        "checks": [it.model_dump(mode="json") for it in items],
    }
    out.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return len(items)


if __name__ == "__main__":
    import argparse
    p = argparse.ArgumentParser()
    p.add_argument("--excel", default="data/source/售前CK.xlsx")
    p.add_argument("--output", default="data/normalized/direct_tender_checks_v3.json")
    args = p.parse_args()
    n = parse_to_file(args.excel, args.output)
    print(f"[OK] wrote {n} checks (v0.3) to {args.output}")